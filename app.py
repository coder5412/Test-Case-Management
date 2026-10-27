import os, sqlite3, csv, io, json, re
from datetime import datetime, timezone
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, make_response, abort, jsonify
from firebase_sync import sync_to_firebase, firebase_status
from openpyxl import load_workbook

BASE=os.path.dirname(os.path.abspath(__file__))
DB=os.getenv('QA_DB_PATH', '/tmp/qa_manual.db' if os.getenv('VERCEL') else os.path.join(BASE,'qa_manual.db'))
app=Flask(__name__)
app.secret_key=os.environ.get('QA_SECRET_KEY','local-only-change-me')
app.config['MAX_CONTENT_LENGTH']=10*1024*1024

def now():return datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
def db():
 c=sqlite3.connect(DB,timeout=20);c.row_factory=sqlite3.Row;c.execute('PRAGMA foreign_keys=ON');return c

def init():
 with db() as c:
  c.executescript('''
  CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL,created_at TEXT NOT NULL);
  CREATE TABLE IF NOT EXISTS cases(id INTEGER PRIMARY KEY,project_id INTEGER NOT NULL,ref TEXT NOT NULL,title TEXT NOT NULL,category TEXT DEFAULT '',steps TEXT DEFAULT '',expected TEXT DEFAULT '',priority TEXT DEFAULT 'Medium',severity TEXT DEFAULT 'Medium',revision INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL,updated_at TEXT NOT NULL,active INTEGER NOT NULL DEFAULT 1,UNIQUE(project_id,ref),FOREIGN KEY(project_id) REFERENCES projects(id));
  CREATE TABLE IF NOT EXISTS case_history(id INTEGER PRIMARY KEY,case_id INTEGER NOT NULL,revision INTEGER NOT NULL,snapshot TEXT NOT NULL,at TEXT NOT NULL,FOREIGN KEY(case_id) REFERENCES cases(id));
  CREATE TABLE IF NOT EXISTS counters(project_id INTEGER PRIMARY KEY,last_version INTEGER NOT NULL DEFAULT 0);
  CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY,project_id INTEGER NOT NULL,version INTEGER NOT NULL,title TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'In Progress',tester TEXT DEFAULT '',environment TEXT DEFAULT '',created_at TEXT NOT NULL,closed_at TEXT,UNIQUE(project_id,version),FOREIGN KEY(project_id) REFERENCES projects(id));
  CREATE TABLE IF NOT EXISTS run_items(id INTEGER PRIMARY KEY,run_id INTEGER NOT NULL,case_id INTEGER,ref TEXT NOT NULL,title TEXT NOT NULL,revision INTEGER NOT NULL,steps TEXT NOT NULL,expected TEXT NOT NULL,priority TEXT NOT NULL,result TEXT NOT NULL DEFAULT 'Not Run',actual TEXT NOT NULL DEFAULT '',notes TEXT NOT NULL DEFAULT '',tested_at TEXT,tested_by TEXT DEFAULT '',UNIQUE(run_id,ref),FOREIGN KEY(run_id) REFERENCES runs(id) ON DELETE CASCADE);
  CREATE TABLE IF NOT EXISTS bugs(id INTEGER PRIMARY KEY,project_id INTEGER NOT NULL,run_id INTEGER,run_item_id INTEGER,case_ref TEXT DEFAULT '',title TEXT NOT NULL,description TEXT DEFAULT '',steps TEXT DEFAULT '',expected TEXT DEFAULT '',actual TEXT DEFAULT '',severity TEXT DEFAULT 'Medium',priority TEXT DEFAULT 'Medium',status TEXT DEFAULT 'Open',assignee TEXT DEFAULT '',reporter TEXT DEFAULT '',environment TEXT DEFAULT '',created_at TEXT NOT NULL,updated_at TEXT NOT NULL,FOREIGN KEY(project_id) REFERENCES projects(id),FOREIGN KEY(run_id) REFERENCES runs(id) ON DELETE SET NULL,FOREIGN KEY(run_item_id) REFERENCES run_items(id) ON DELETE SET NULL);
  CREATE TABLE IF NOT EXISTS bug_comments(id INTEGER PRIMARY KEY,bug_id INTEGER NOT NULL,author TEXT NOT NULL,message TEXT NOT NULL,created_at TEXT NOT NULL,FOREIGN KEY(bug_id) REFERENCES bugs(id) ON DELETE CASCADE);
  CREATE TABLE IF NOT EXISTS bug_events(id INTEGER PRIMARY KEY,bug_id INTEGER NOT NULL,action TEXT NOT NULL,at TEXT NOT NULL,FOREIGN KEY(bug_id) REFERENCES bugs(id) ON DELETE CASCADE);
  ''')
  c.execute('INSERT OR IGNORE INTO projects(name,created_at) VALUES(?,?)',('Default Project',now()))

ALLOWED_RESULTS={'Not Run','Passed','Failed','Blocked','Skipped'}
ALLOWED_BUG={'Open','In Progress','Ready for Retest','Resolved','Closed','Reopened','Rejected'}
FIELDS=['title','category','steps','expected','priority','severity']
def snap(row):return json.dumps({k:row[k] for k in ['ref']+FIELDS},ensure_ascii=False)
def project_rows(c):return c.execute('SELECT * FROM projects ORDER BY name').fetchall()
def get_project(c,pid):
 p=c.execute('SELECT * FROM projects WHERE id=?',(pid,)).fetchone()
 if not p:abort(404)
 return p
@app.context_processor
def globals_():
 with db() as c:
  return {'projects':project_rows(c),'datetime_now':now()}

def page(view,**kwargs):return render_template('index.html',view=view,**kwargs)
@app.route('/')
def dashboard():
 with db() as c:
  stats={k:c.execute(q).fetchone()[0] for k,q in {'cases':'SELECT count(*) FROM cases WHERE active=1','bugs':'SELECT count(*) FROM bugs WHERE status NOT IN ("Closed","Rejected")','runs':'SELECT count(*) FROM runs','projects':'SELECT count(*) FROM projects'}.items()}
  recent=c.execute('SELECT r.*,p.name project FROM runs r JOIN projects p ON p.id=r.project_id ORDER BY r.id DESC LIMIT 8').fetchall()
  active=c.execute('SELECT * FROM bugs ORDER BY id DESC LIMIT 6').fetchall()
 return page('dashboard',stats=stats,recent=recent,active=active)

@app.post('/projects')
def new_project():
 name=request.form.get('name','').strip()
 if name:
  with db() as c:
   try:c.execute('INSERT INTO projects(name,created_at) VALUES(?,?)',(name,now()));flash('Project created.','ok')
   except sqlite3.IntegrityError:flash('Project already exists.','warn')
 return redirect(request.referrer or url_for('cases'))

@app.post('/projects/quick-create')
def project_quick_create():
 name=request.form.get('name','').strip()
 if not name or len(name)>100:
  flash('Please enter a project name up to 100 characters.','warn')
  return redirect(url_for('case_new'))
 with db() as c:
  found=c.execute('SELECT id FROM projects WHERE name=? COLLATE NOCASE',(name,)).fetchone()
  if found:
   pid=found['id']
   flash('Existing project selected.','ok')
  else:
   try:
    cur=c.execute('INSERT INTO projects(name,created_at) VALUES(?,?)',(name,now()))
    pid=cur.lastrowid
    flash('New project created and selected.','ok')
   except sqlite3.IntegrityError:
    flash('Project already exists.','warn')
    return redirect(url_for('case_new'))
 return redirect(url_for('case_new',project_id=pid))

@app.post('/projects/<int:pid>/rename')
def rename_project(pid):
 name=request.form.get('name','').strip()
 if not name or len(name)>100:
  flash('Enter a valid project name (maximum 100 characters).','warn')
  return redirect(url_for('settings'))
 with db() as c:
  get_project(c,pid)
  existing=c.execute('SELECT id FROM projects WHERE name=? COLLATE NOCASE AND id!=?',(name,pid)).fetchone()
  if existing:flash('Another project already has this name.','warn')
  else:
   try:c.execute('UPDATE projects SET name=? WHERE id=?',(name,pid));flash('Project renamed. Linked cases, runs and bugs remain intact.','ok')
   except sqlite3.IntegrityError:flash('Project name is already in use.','warn')
 return redirect(url_for('settings'))

@app.route('/cases')
def cases():
 pid=request.args.get('project',type=int);q=request.args.get('q','').strip()
 with db() as c:
  sql='SELECT ca.*,p.name project FROM cases ca JOIN projects p ON ca.project_id=p.id WHERE ca.active=1';params=[]
  if pid:sql+=' AND ca.project_id=?';params.append(pid)
  if q:sql+=' AND (ca.ref LIKE ? OR ca.title LIKE ? OR ca.category LIKE ?)';params += ['%'+q+'%']*3
  sql+=' ORDER BY ca.id DESC'
  rows=c.execute(sql,params).fetchall()
 return page('cases',cases=rows,chosen=pid,q=q)

@app.route('/cases/new',methods=['GET','POST'])
def case_new():
 if request.method=='POST':
  f=request.form;pid=int(f['project_id']);ref=f.get('ref','').strip();title=f.get('title','').strip()
  if not ref or not title:flash('Case ID and title required.','warn');return redirect(request.url)
  vals=[f.get(k,'').strip() for k in FIELDS]
  with db() as c:
   try:
    cur=c.execute('INSERT INTO cases(project_id,ref,title,category,steps,expected,priority,severity,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(pid,ref,*vals,now(),now()))
    row=c.execute('SELECT * FROM cases WHERE id=?',(cur.lastrowid,)).fetchone()
    c.execute('INSERT INTO case_history(case_id,revision,snapshot,at) VALUES(?,?,?,?)',(cur.lastrowid,1,snap(row),now()))
    flash('Test case created (revision 1).','ok')
   except sqlite3.IntegrityError:flash('Test ID already exists for that project.','warn')
  return redirect(url_for('cases',project=pid))
 return page('case_form',case=None,hist=[],selected_project=request.args.get('project_id',type=int))

@app.route('/cases/<int:cid>/edit',methods=['GET','POST'])
def case_edit(cid):
 with db() as c:
  case=c.execute('SELECT * FROM cases WHERE id=?',(cid,)).fetchone()
  if not case:abort(404)
  if request.method=='POST':
   vals=[request.form.get(k,'').strip() for k in FIELDS]
   if not vals[0]:flash('Title required.','warn');return redirect(request.url)
   if any((case[k] or '')!=v for k,v in zip(FIELDS,vals)):
    rev=case['revision']+1
    c.execute('UPDATE cases SET title=?,category=?,steps=?,expected=?,priority=?,severity=?,revision=?,updated_at=? WHERE id=?',(*vals,rev,now(),cid))
    updated=c.execute('SELECT * FROM cases WHERE id=?',(cid,)).fetchone()
    c.execute('INSERT INTO case_history(case_id,revision,snapshot,at) VALUES(?,?,?,?)',(cid,rev,snap(updated),now()))
    flash(f'Saved as revision {rev}.','ok')
   else:flash('No changes; revision remains unchanged.','ok')
   return redirect(url_for('case_edit',cid=cid))
  hist=c.execute('SELECT * FROM case_history WHERE case_id=? ORDER BY revision DESC',(cid,)).fetchall()
 return page('case_form',case=case,hist=hist)

@app.post('/cases/<int:cid>/archive')
def archive_case(cid):
 with db() as c:c.execute('UPDATE cases SET active=0 WHERE id=?',(cid,))
 flash('Case archived. Existing run snapshots preserved.','ok')
 return redirect(url_for('cases'))

def colkey(s):return re.sub(r'[^a-z0-9]+',' ',str(s or '').lower()).strip()
def pick(row,*keys):
 for k in keys:
  for col,val in row.items():
   if colkey(col)==colkey(k) and val is not None:return str(val).strip()
 return ''
@app.route('/import',methods=['GET','POST'])
def import_excel():
 if request.method=='POST':
  f=request.files.get('file');pid=request.form.get('project_id',type=int)
  if not f or not f.filename.lower().endswith('.xlsx'):flash('Upload .xlsx file.','warn');return redirect(request.url)
  try:
   wb=load_workbook(f.stream,read_only=True,data_only=True)
   added=updated=skipped=0
   with db() as c:
    get_project(c,pid)
    for ws in wb.worksheets:
     iterator=ws.iter_rows(values_only=True)
     names=[]
     for candidate in iterator:
      labels=[colkey(x) for x in candidate]
      if any(x in ('test id','test case id','case id') for x in labels) and any(x in ('test scenario','test case','title','scenario') for x in labels):
       names=[str(x or '').strip() for x in candidate]
       break
     if not names:continue
     for cells in iterator:
      row=dict(zip(names,cells))
      ref=pick(row,'Test ID','Test Case ID','Case ID','ID','Test Case No')
      title=pick(row,'Test Scenario','Test Case','Title','Scenario','Test Case Description')
      if not ref or not title:skipped+=1;continue
      vals=[title,pick(row,'Category','Module'),pick(row,'Test Steps','Steps','Steps to Reproduce'),pick(row,'Expected Result','Expected'),pick(row,'Priority') or 'Medium',pick(row,'Severity') or 'Medium']
      old=c.execute('SELECT * FROM cases WHERE project_id=? AND ref=?',(pid,ref)).fetchone()
      if old:
       changed=any((old[k] or '')!=v for k,v in zip(FIELDS,vals))
       if changed:
        rev=old['revision']+1
        c.execute('UPDATE cases SET title=?,category=?,steps=?,expected=?,priority=?,severity=?,revision=?,updated_at=?,active=1 WHERE id=?',(*vals,rev,now(),old['id']))
        revised=c.execute('SELECT * FROM cases WHERE id=?',(old['id'],)).fetchone()
        c.execute('INSERT INTO case_history(case_id,revision,snapshot,at) VALUES(?,?,?,?)',(old['id'],rev,snap(revised),now()))
        updated+=1
       else:skipped+=1
      else:
       cur=c.execute('INSERT INTO cases(project_id,ref,title,category,steps,expected,priority,severity,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(pid,ref,*vals,now(),now()))
       added+=1
       new=c.execute('SELECT * FROM cases WHERE id=?',(cur.lastrowid,)).fetchone()
       c.execute('INSERT INTO case_history(case_id,revision,snapshot,at) VALUES(?,?,?,?)',(cur.lastrowid,1,snap(new),now()))
   flash(f'Import complete: {added} new, {updated} revised, {skipped} unchanged/skipped.','ok')
   return redirect(url_for('cases',project=pid))
  except Exception as e:flash('Import failed: '+str(e)[:200],'warn')
 return page('import')

@app.route('/runs')
def runs():
 with db() as c:
  rows=c.execute('''SELECT r.*,p.name project,(SELECT COUNT(*) FROM run_items ri WHERE ri.run_id=r.id) total,(SELECT COUNT(*) FROM run_items ri WHERE ri.run_id=r.id AND result='Passed') passed,(SELECT COUNT(*) FROM run_items ri WHERE ri.run_id=r.id AND result='Failed') failed FROM runs r JOIN projects p ON p.id=r.project_id ORDER BY r.id DESC''').fetchall()
 return page('runs',runs=rows)

@app.post('/runs/new')
def start_run():
 f=request.form;pid=int(f['project_id']);title=f.get('title','').strip() or 'Manual test cycle'
 with db() as c:
  get_project(c,pid)
  allcases=c.execute('SELECT * FROM cases WHERE project_id=? AND active=1 ORDER BY ref',(pid,)).fetchall()
  if not allcases:flash('Import or create test cases first.','warn');return redirect(url_for('runs'))
  c.execute('INSERT OR IGNORE INTO counters(project_id,last_version) VALUES(?,COALESCE((SELECT MAX(version) FROM runs WHERE project_id=?),0))',(pid,pid))
  c.execute('UPDATE counters SET last_version=MAX(last_version,COALESCE((SELECT MAX(version) FROM runs WHERE project_id=?),0))+1 WHERE project_id=?',(pid,pid))
  ver=c.execute('SELECT last_version FROM counters WHERE project_id=?',(pid,)).fetchone()[0]
  cur=c.execute('INSERT INTO runs(project_id,version,title,tester,environment,created_at) VALUES(?,?,?,?,?,?)',(pid,ver,title,f.get('tester','').strip(),f.get('environment','').strip(),now()))
  rid=cur.lastrowid
  c.executemany('INSERT INTO run_items(run_id,case_id,ref,title,revision,steps,expected,priority) VALUES(?,?,?,?,?,?,?,?)',[(rid,x['id'],x['ref'],x['title'],x['revision'],x['steps'],x['expected'],x['priority']) for x in allcases])
 flash(f'Execution V{ver} created with {len(allcases)} cases.','ok')
 return redirect(url_for('run_detail',rid=rid))

@app.route('/runs/<int:rid>')
def run_detail(rid):
 with db() as c:
  run=c.execute('SELECT r.*,p.name project FROM runs r JOIN projects p ON r.project_id=p.id WHERE r.id=?',(rid,)).fetchone()
  if not run:abort(404)
  items=c.execute('SELECT * FROM run_items WHERE run_id=? ORDER BY ref',(rid,)).fetchall()
  bugs=c.execute('SELECT id,title,status,case_ref FROM bugs WHERE run_id=? ORDER BY id DESC',(rid,)).fetchall()
 return page('run_detail',run=run,items=items,related_bugs=bugs)

@app.post('/items/<int:iid>')
def update_item(iid):
 f=request.form
 with db() as c:
  row=c.execute('SELECT ri.*,r.status run_status FROM run_items ri JOIN runs r ON r.id=ri.run_id WHERE ri.id=?',(iid,)).fetchone()
  if not row:abort(404)
  if row['run_status']=='Closed':flash('Closed run is read-only.','warn');return redirect(url_for('run_detail',rid=row['run_id']))
  result=f.get('result','Not Run')
  if result not in ALLOWED_RESULTS:abort(400)
  c.execute('UPDATE run_items SET result=?,actual=?,notes=?,tested_by=?,tested_at=? WHERE id=?',(result,f.get('actual','').strip(),f.get('notes','').strip(),f.get('tested_by','').strip(),now(),iid))
 flash('Result saved in this execution version.','ok')
 return redirect(url_for('run_detail',rid=row['run_id'])+'#item-'+str(iid))

@app.post('/runs/<int:rid>/close')
def close_run(rid):
 with db() as c:c.execute("UPDATE runs SET status='Closed',closed_at=? WHERE id=? AND status!='Closed'",(now(),rid))
 flash('Run closed and locked.','ok');return redirect(url_for('run_detail',rid=rid))

@app.post('/runs/<int:rid>/delete')
def delete_run(rid):
 with db() as c:
  c.execute('UPDATE bugs SET run_item_id=NULL,run_id=NULL WHERE run_id=?',(rid,))
  c.execute('DELETE FROM run_items WHERE run_id=?',(rid,));c.execute('DELETE FROM runs WHERE id=?',(rid,))
 flash('Version deleted. Next execution will use the next higher version number; linked bugs preserved.','ok')
 return redirect(url_for('runs'))

@app.route('/bugs')
def bugs():
 with db() as c:rows=c.execute('SELECT b.*,p.name project FROM bugs b JOIN projects p ON p.id=b.project_id ORDER BY b.id DESC').fetchall()
 return page('bugs',bugs=rows)

@app.route('/bugs/new',methods=['GET','POST'])
def bug_new():
 if request.method=='POST':
  f=request.form;pid=int(f['project_id']);rid=f.get('run_id',type=int);iid=f.get('run_item_id',type=int)
  with db() as c:
   get_project(c,pid)
   if iid:
    valid=c.execute('SELECT ri.* FROM run_items ri JOIN runs r ON r.id=ri.run_id WHERE ri.id=? AND r.project_id=?',(iid,pid)).fetchone()
    if not valid:abort(400)
    rid=valid['run_id']
   cur=c.execute('''INSERT INTO bugs(project_id,run_id,run_item_id,case_ref,title,description,steps,expected,actual,severity,priority,status,assignee,reporter,environment,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(pid,rid,iid,f.get('case_ref','').strip(),f.get('title','').strip(),f.get('description','').strip(),f.get('steps','').strip(),f.get('expected','').strip(),f.get('actual','').strip(),f.get('severity','Medium'),f.get('priority','Medium'),'Open',f.get('assignee','').strip(),f.get('reporter','').strip(),f.get('environment','').strip(),now(),now()))
   c.execute('INSERT INTO bug_events(bug_id,action,at) VALUES(?,?,?)',(cur.lastrowid,'Bug created',now()))
  flash('Bug created.','ok');return redirect(url_for('bug_detail',bid=cur.lastrowid))
  
 with db() as c:
  pid=request.args.get('project_id',type=int);item_id=request.args.get('item_id',type=int)
  pre={}
  if item_id:
   item=c.execute('SELECT ri.*,r.project_id,r.environment FROM run_items ri JOIN runs r ON r.id=ri.run_id WHERE ri.id=?',(item_id,)).fetchone()
   if item:pre=dict(item);pid=item['project_id']
 return page('bug_form',bug=None,pre=pre,pref_project=pid,comments=[],events=[])

@app.route('/bugs/<int:bid>',methods=['GET','POST'])
def bug_detail(bid):
 with db() as c:
  bug=c.execute('SELECT * FROM bugs WHERE id=?',(bid,)).fetchone()
  if not bug:abort(404)
  if request.method=='POST':
   f=request.form;status=f.get('status',bug['status'])
   if status not in ALLOWED_BUG:abort(400)
   cols=['title','description','steps','expected','actual','severity','priority','assignee','reporter','environment','case_ref']
   vals=[f.get(k,'').strip() for k in cols]
   c.execute('UPDATE bugs SET '+','.join(k+'=?' for k in cols)+',status=?,updated_at=? WHERE id=?',(*vals,status,now(),bid))
   c.execute('INSERT INTO bug_events(bug_id,action,at) VALUES(?,?,?)',(bid,'Bug updated / '+status,now()))
   flash('Bug updated.','ok');return redirect(url_for('bug_detail',bid=bid))
  comments=c.execute('SELECT * FROM bug_comments WHERE bug_id=? ORDER BY id DESC',(bid,)).fetchall()
  events=c.execute('SELECT * FROM bug_events WHERE bug_id=? ORDER BY id DESC',(bid,)).fetchall()
 return page('bug_form',bug=bug,pre={},pref_project=bug['project_id'],comments=comments,events=events)

@app.post('/bugs/<int:bid>/comment')
def bug_comment(bid):
 f=request.form;msg=f.get('message','').strip()
 if msg:
  with db() as c:
   c.execute('INSERT INTO bug_comments(bug_id,author,message,created_at) VALUES(?,?,?,?)',(bid,f.get('author','QA Tester').strip(),msg,now()))
   c.execute('INSERT INTO bug_events(bug_id,action,at) VALUES(?,?,?)',(bid,'Comment added',now()))
 return redirect(url_for('bug_detail',bid=bid))

@app.route('/reports')
def reports():
 with db() as c:
  runs=c.execute('SELECT r.*,p.name project FROM runs r JOIN projects p ON p.id=r.project_id ORDER BY r.id DESC').fetchall()
 return page('reports',runs=runs)

@app.route('/runs/<int:rid>/export.csv')
def export_csv(rid):
 with db() as c:
  run=c.execute('SELECT r.*,p.name project FROM runs r JOIN projects p ON p.id=r.project_id WHERE r.id=?',(rid,)).fetchone()
  if not run:abort(404)
  rows=c.execute('SELECT * FROM run_items WHERE run_id=? ORDER BY ref',(rid,)).fetchall()
 s=io.StringIO();w=csv.writer(s);w.writerow(['Project','Execution Version','Test ID','Scenario','Case Revision','Test Steps','Expected Result','Result','Actual Result','Notes','Tester','Tested At'])
 for r in rows:w.writerow([run['project'],'V'+str(run['version']),r['ref'],r['title'],r['revision'],r['steps'],r['expected'],r['result'],r['actual'],r['notes'],r['tested_by'],r['tested_at']])
 resp=make_response('\ufeff'+s.getvalue());resp.headers['Content-Type']='text/csv; charset=utf-8';resp.headers['Content-Disposition']=f'attachment; filename=QA_Execution_V{run["version"]}.csv';return resp

@app.route('/runs/compare')
def compare():
 a=request.args.get('a',type=int);b=request.args.get('b',type=int)
 with db() as c:
  runs=c.execute('SELECT r.*,p.name project FROM runs r JOIN projects p ON p.id=r.project_id ORDER BY r.id DESC').fetchall()
  ra=c.execute('SELECT * FROM runs WHERE id=?',(a,)).fetchone() if a else None
  rb=c.execute('SELECT * FROM runs WHERE id=?',(b,)).fetchone() if b else None
  results=[]
  if ra and rb and ra['project_id']==rb['project_id']:
   ia={r['ref']:r for r in c.execute('SELECT * FROM run_items WHERE run_id=?',(a,))}
   ib={r['ref']:r for r in c.execute('SELECT * FROM run_items WHERE run_id=?',(b,))}
   results=[{'ref':key,'title':(ib.get(key) or ia.get(key))['title'],'a':ia[key]['result'] if key in ia else '—','b':ib[key]['result'] if key in ib else '—'} for key in sorted(set(ia)|set(ib))]
 return page('compare',runs=runs,ra=ra,rb=rb,comparisons=results)

@app.route('/settings')
def settings():return page('settings',firebase=firebase_status())

@app.post('/firebase/sync')
def firebase_sync_now():
 result=sync_to_firebase(DB)
 flash('Firebase sync completed.' if result['ok'] else 'Firebase sync failed: '+result['message'][:240], 'ok' if result['ok'] else 'warn')
 return redirect(url_for('settings'))

@app.get('/api/firebase/status')
def firebase_status_api():
 return jsonify(firebase_status())

@app.after_request
def firebase_mirror_after_write(response):
 if request.method in ('POST','PUT','PATCH','DELETE') and response.status_code < 400 and request.path != '/firebase/sync':
  # SQLite transaction is committed when the route exits its database context.
  # Firebase failure cannot undo successful local edits.
  sync_to_firebase(DB)
 return response

@app.route('/health')
def health():return jsonify({'status':'ok','mode':'manual-test-management'})

if os.getenv('VERCEL'):
 init()

if __name__=='__main__':
 init();app.run(host='127.0.0.1',port=int(os.environ.get('PORT','5000')),debug=False)
